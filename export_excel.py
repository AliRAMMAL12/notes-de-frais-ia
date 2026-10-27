"""Export Excel propre de la note de frais (en-tête, largeurs, totaux)."""

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NOM_FEUILLE = "Notes de frais"


def exporter_excel(df: pd.DataFrame) -> bytes:
    sortie = BytesIO()

    with pd.ExcelWriter(sortie, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=NOM_FEUILLE)
        ws = writer.sheets[NOM_FEUILLE]

        # En-tête
        for cellule in ws[1]:
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor="2F5D50")
            cellule.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Largeur des colonnes selon le contenu
        for i, colonne in enumerate(df.columns, start=1):
            largeur = max([len(str(colonne))] + [len(str(v)) for v in df[colonne]])
            ws.column_dimensions[get_column_letter(i)].width = min(largeur + 3, 60)

        # Colonnes en euros : format monétaire + ligne de total (formule SUM)
        nb = len(df)
        ligne_total = nb + 2
        for i, colonne in enumerate(df.columns, start=1):
            if "(€)" not in colonne:
                continue
            lettre = get_column_letter(i)
            for ligne in range(2, nb + 2):
                ws.cell(row=ligne, column=i).number_format = '#,##0.00 "€"'
            total = ws.cell(row=ligne_total, column=i, value=f"=SUM({lettre}2:{lettre}{nb + 1})")
            total.number_format = '#,##0.00 "€"'
            total.font = Font(bold=True)

        ws.cell(row=ligne_total, column=1, value="TOTAL").font = Font(bold=True)

    return sortie.getvalue()
